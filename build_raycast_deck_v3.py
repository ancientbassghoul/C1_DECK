from pathlib import Path
import argparse

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
DEFAULT_OUTPUT = ROOT / "presentation_v3_design_proof.pptx"
SLIDE_1_VISUAL = ROOT / "visuals" / "Raycast_Slide_1_Visual.png"
SLIDE_2_VISUAL = ROOT / "visuals" / "Raycast_Slide_2_Visual.png"
SLIDE_3_SHARP = Path(r"D:\EXTEND\C1_IMGS\2026-02-15_16-25-03_12035.png")
SLIDE_3_BLURRED = Path(r"D:\EXTEND\C1_IMGS\2026-02-15_16-25-03_10305.png")
SLIDE_5_VIDEO = ROOT / "visuals" / "Raycast_Slide_8_Visual_RETIMED.mp4"
SLIDE_5_POSTER = ROOT / "visuals" / "Raycast_Slide_8_Poster.jpg"
SLIDE_6_VISUAL = ROOT / "visuals" / "Raycast_Slide_9_Visual.png"
SLIDE_7_VIDEO = ROOT / "visuals" / "Raycast_Slide_10_Visual.mp4"
SLIDE_7_POSTER = ROOT / "visuals" / "Raycast_Slide_10_Poster.jpg"
SLIDE_8_VISUAL = ROOT / "visuals" / "Raycast_Slide_11_Visual.png"
SLIDE_9_VISUAL = ROOT / "visuals" / "Raycast_Slide_20_NEW_Visual.png"
SLIDE_10_VISUAL = ROOT / "visuals" / "Raycast_Slide_21_Visual.png"
SLIDE_11_VISUAL = ROOT / "visuals" / "Raycast_Slide_23_Visual.png"
SLIDE_12_VISUAL = ROOT / "visuals" / "Raycast_Slide_23.5_Visual.png"
SLIDE_13_VISUAL = ROOT / "visuals" / "Raycast_Slide_24_Visual.png"
SLIDE_14_VISUAL = ROOT / "visuals" / "Raycast_Slide_28_Poster.jpg"
SLIDE_15_BLACK = ROOT / "visuals" / "Raycast_Slide_29_01_Visual.png"
SLIDE_15_FEATHERED = ROOT / "visuals" / "Raycast_Slide_29_02_Visual.png"
SLIDE_16_VISUAL = ROOT / "visuals" / "Raycast_Slide_32_Visual.png"
SLIDE_17_VIDEO = ROOT / "visuals" / "Raycast_Slide_35_Visual.mp4"
SLIDE_17_POSTER = ROOT / "visuals" / "Raycast_Slide_35_Poster.jpg"
SLIDE_18_VISUAL = ROOT / "results_slide" / "How_we_measure_Visual.png"
RESULTS_MATRIX = ROOT / "results_slide" / "score_pairs_matrix.png"
APPENDIX_A1_VISUAL = ROOT / "visuals" / "Raycast_Slide_33_Visual.png"

SLIDE_W = 13.333
SLIDE_H = 7.5

MIDNIGHT = "101827"
DEEP = "09111E"
GREEN = "27E38D"
CORAL = "FF5B5B"
OFFWHITE = "F5F2EA"
WHITE = "FFFFFF"
INK = "172033"
MUTED = "687386"
PALE = "DCE2EA"
LIGHT_PANEL = "E8EDF2"

HEADER_FONT = "Bahnschrift SemiBold"
BODY_FONT = "Aptos"
CODE_FONT = "Cascadia Mono"

WARNINGS = []


def rgb(value):
    return RGBColor.from_string(value)


def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def set_shape_alpha(shape, opacity_percent):
    """Set solid-fill opacity with DrawingML; 100 is opaque, 0 is invisible."""
    solid_fill = shape.fill._xPr.solidFill
    color_node = solid_fill[0]
    for old_alpha in color_node.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}alpha"):
        color_node.remove(old_alpha)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(int(opacity_percent * 1000)))
    color_node.append(alpha)


def add_rect(
    slide,
    x,
    y,
    w,
    h,
    fill,
    line=None,
    radius=False,
    line_width=1.0,
    opacity=100,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if opacity != 100:
        set_shape_alpha(shape, opacity)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_line(slide, x1, y1, x2, y2, color, width=1.0, dash=None):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x1),
        Inches(y1),
        Inches(max(x2 - x1, 0.008)),
        Inches(max(y2 - y1, 0.008)),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(color)
    line.line.fill.background()
    return line


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size,
    color,
    font=BODY_FONT,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    italic=False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return box


def add_rich_line(slide, segments, x, y, w, h, size, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    for segment in segments:
        run = paragraph.add_run()
        run.text = segment["text"]
        run.font.name = segment.get("font", BODY_FONT)
        run.font.size = Pt(segment.get("size", size))
        run.font.bold = segment.get("bold", False)
        run.font.italic = segment.get("italic", False)
        run.font.color.rgb = rgb(segment.get("color", color))
    return box


def set_text_outline(textbox, color, width_pt):
    """Apply a solid outline to every run in a textbox."""
    width_emu = str(int(width_pt * 12700))
    for paragraph in textbox.text_frame.paragraphs:
        for run in paragraph.runs:
            run_properties = run._r.get_or_add_rPr()
            existing = run_properties.find(qn("a:ln"))
            if existing is not None:
                run_properties.remove(existing)
            line = OxmlElement("a:ln")
            line.set("w", width_emu)
            solid_fill = OxmlElement("a:solidFill")
            srgb = OxmlElement("a:srgbClr")
            srgb.set("val", color)
            solid_fill.append(srgb)
            line.append(solid_fill)
            dash = OxmlElement("a:prstDash")
            dash.set("val", "solid")
            line.append(dash)
            run_properties.append(line)


def add_picture_cover(slide, image_path, x, y, w, h):
    with Image.open(image_path) as image:
        image_w, image_h = image.size
    image_ratio = image_w / image_h
    frame_ratio = w / h
    picture = slide.shapes.add_picture(
        str(image_path), Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if image_ratio > frame_ratio:
        visible_fraction = frame_ratio / image_ratio
        crop = (1 - visible_fraction) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    elif image_ratio < frame_ratio:
        visible_fraction = image_ratio / frame_ratio
        crop = (1 - visible_fraction) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def add_picture_contain(slide, image_path, x, y, w, h):
    with Image.open(image_path) as image:
        image_w, image_h = image.size
    scale = min(w / image_w, h / image_h)
    placed_w = image_w * scale
    placed_h = image_h * scale
    placed_x = x + (w - placed_w) / 2
    placed_y = y + (h - placed_h) / 2
    return slide.shapes.add_picture(
        str(image_path),
        Inches(placed_x),
        Inches(placed_y),
        Inches(placed_w),
        Inches(placed_h),
    )


def add_missing_asset(slide, label, x, y, w, h, dark=False):
    """Place an explicit, presentation-safe placeholder and report the missing path."""
    WARNINGS.append(f"Missing asset: {label}")
    fill = "202B3C" if dark else LIGHT_PANEL
    line = CORAL
    text_color = WHITE if dark else INK
    add_rect(slide, x, y, w, h, fill, line=line, radius=True, line_width=1.5)
    add_text(
        slide,
        f"MISSING ASSET\n{label}",
        x + 0.25,
        y + 0.25,
        w - 0.5,
        h - 0.5,
        16,
        text_color,
        font=CODE_FONT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_movie_contain(slide, movie_path, poster_path, x, y, w, h, shape_name, dark=False):
    """Embed a video without distorting its poster or playback frame."""
    if not movie_path.exists() or not poster_path.exists():
        missing = movie_path if not movie_path.exists() else poster_path
        add_missing_asset(slide, str(missing), x, y, w, h, dark=dark)
        return None

    with Image.open(poster_path) as image:
        image_w, image_h = image.size
    scale = min(w / image_w, h / image_h)
    placed_w = image_w * scale
    placed_h = image_h * scale
    placed_x = x + (w - placed_w) / 2
    placed_y = y + (h - placed_h) / 2
    movie = slide.shapes.add_movie(
        str(movie_path),
        Inches(placed_x),
        Inches(placed_y),
        Inches(placed_w),
        Inches(placed_h),
        poster_frame_image=str(poster_path),
        mime_type="video/mp4",
    )
    movie.name = shape_name
    return movie


def configure_video_click_timing(
    prs, slide_number, shape_name, duration_ms, fade_shape_names=None, fade_duration_ms=350
):
    """Play an embedded video on click, optionally fading overlays with it."""
    if len(prs.slides) < slide_number:
        return
    slide = prs.slides[slide_number - 1]
    video = next((shape for shape in slide.shapes if shape.name == shape_name), None)
    if video is None:
        return

    slide_element = slide._element
    existing_timing = slide_element.find(qn("p:timing"))
    if existing_timing is not None:
        slide_element.remove(existing_timing)

    shape_id = video.shape_id
    fade_effects = []
    build_entries = []
    next_timing_id = 8
    for index, fade_name in enumerate(fade_shape_names or []):
        fade_shape = next(
            (shape for shape in slide.shapes if shape.name == fade_name), None
        )
        if fade_shape is None:
            continue
        fade_shape_id = fade_shape.shape_id
        effect_id = next_timing_id
        animation_id = next_timing_id + 1
        visibility_id = next_timing_id + 2
        next_timing_id += 3
        fade_effects.append(
            f"""
                                    <p:par>
                                      <p:cTn id="{effect_id}" presetID="10" presetClass="exit" presetSubtype="0" fill="hold" grpId="0" nodeType="withEffect">
                                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                                        <p:childTnLst>
                                          <p:animEffect transition="out" filter="fade">
                                            <p:cBhvr>
                                              <p:cTn id="{animation_id}" dur="{fade_duration_ms}"/>
                                              <p:tgtEl><p:spTgt spid="{fade_shape_id}"/></p:tgtEl>
                                            </p:cBhvr>
                                          </p:animEffect>
                                          <p:set>
                                            <p:cBhvr>
                                              <p:cTn id="{visibility_id}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{fade_duration_ms - 1}"/></p:stCondLst></p:cTn>
                                              <p:tgtEl><p:spTgt spid="{fade_shape_id}"/></p:tgtEl>
                                              <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                                            </p:cBhvr>
                                            <p:to><p:strVal val="hidden"/></p:to>
                                          </p:set>
                                        </p:childTnLst>
                                      </p:cTn>
                                    </p:par>
            """
        )
        animation_background = ' animBg="1"' if index == 0 else ""
        build_entries.append(
            f'<p:bldP spid="{fade_shape_id}" grpId="0"{animation_background}/>'
        )

    fade_effects_xml = "".join(fade_effects)
    build_list_xml = (
        f"<p:bldLst>{''.join(build_entries)}</p:bldLst>" if build_entries else ""
    )
    timing = parse_xml(
        f"""
        <p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
          <p:tnLst>
            <p:par>
              <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
                <p:childTnLst>
                  <p:seq concurrent="1" nextAc="seek">
                    <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                      <p:childTnLst>
                        <p:par>
                          <p:cTn id="3" fill="hold">
                            <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
                            <p:childTnLst>
                              <p:par>
                                <p:cTn id="4" fill="hold">
                                  <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                                  <p:childTnLst>
                                    <p:par>
                                      <p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="clickEffect">
                                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                                        <p:childTnLst>
                                          <p:cmd type="call" cmd="playFrom(0.0)">
                                            <p:cBhvr>
                                              <p:cTn id="6" dur="{duration_ms}" fill="hold"/>
                                              <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
                                            </p:cBhvr>
                                          </p:cmd>
                                        </p:childTnLst>
                                      </p:cTn>
                                    </p:par>
                                    {fade_effects_xml}
                                  </p:childTnLst>
                                </p:cTn>
                              </p:par>
                            </p:childTnLst>
                          </p:cTn>
                        </p:par>
                      </p:childTnLst>
                    </p:cTn>
                    <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
                    <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
                  </p:seq>
                  <p:video fullScrn="false">
                    <p:cMediaNode vol="80000">
                      <p:cTn id="7" fill="hold" display="0">
                        <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
                      </p:cTn>
                      <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
                    </p:cMediaNode>
                  </p:video>
                </p:childTnLst>
              </p:cTn>
            </p:par>
          </p:tnLst>
          {build_list_xml}
        </p:timing>
        """
    )
    slide_element.append(timing)


def add_slide_title(slide, title, kicker=None, dark=False):
    title_color = WHITE if dark else INK
    kicker_color = GREEN if dark else MUTED
    add_text(
        slide,
        title,
        0.72,
        0.42,
        11.85,
        0.62,
        38,
        title_color,
        font=HEADER_FONT,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    if kicker:
        add_text(
            slide,
            kicker.upper(),
            0.75,
            1.08,
            8.5,
            0.28,
            11,
            kicker_color,
            font=CODE_FONT,
            bold=True,
        )
    add_rect(slide, 0.74, 1.39, 11.85, 0.018, PALE if not dark else "344054")


def add_notes(slide, text):
    notes_frame = slide.notes_slide.notes_text_frame
    if notes_frame is not None:
        notes_frame.text = text


def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_picture_cover(slide, SLIDE_1_VISUAL, 0, 0, SLIDE_W, SLIDE_H)
    add_rect(slide, 1.58, 0.32, 8.36, 1.68, DEEP, radius=True, opacity=75)
    title = add_text(
        slide,
        "The Raycast Challenge",
        1.84,
        0.53,
        7.80,
        0.65,
        50,
        WHITE,
        font=HEADER_FONT,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    set_text_outline(title, "000000", 1.0)
    subtitle = add_text(
        slide,
        "Pick a pixel in one drone frame. Find the same point in twelve others in ≤10 pixels accuracy.",
        1.86,
        1.27,
        7.82,
        0.48,
        19.5,
        OFFWHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )
    set_text_outline(subtitle, "000000", 0.25)


def add_assignment_row(slide, y, number, text, size=18.5):
    add_rect(slide, 5.28, y, 7.32, 1.10, WHITE, line=PALE, radius=True, line_width=1.0)
    add_text(
        slide,
        number,
        5.58,
        y + 0.14,
        0.48,
        0.28,
        11,
        GREEN,
        font=CODE_FONT,
        bold=True,
    )
    add_text(
        slide,
        text,
        6.20,
        y + 0.14,
        6.04,
        0.82,
        size,
        INK,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )


def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_slide_title(slide, "The Assignment", "One sequence · one target · proof required")

    add_rect(slide, 0.72, 1.70, 4.15, 5.12, WHITE, line=PALE, radius=True, line_width=1.0)
    if SLIDE_2_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_2_VISUAL, 0.84, 1.82, 3.91, 4.88)
        picture.name = "Assignment_Visual"
    else:
        add_missing_asset(slide, str(SLIDE_2_VISUAL), 0.84, 1.82, 3.91, 4.88)

    add_assignment_row(
        slide,
        1.70,
        "01",
        "A sequence of drone FPV frames, one stationary van, one agricultural field",
    )
    add_assignment_row(
        slide,
        2.94,
        "02",
        "Several frames badly motion/defocus-blurred",
    )
    add_assignment_row(
        slide,
        4.18,
        "03",
        "Deliverable: source code, a methodology writeup, and proof — not just a claim — of ≤10px reprojection accuracy",
        size=18,
    )
    add_assignment_row(
        slide,
        5.42,
        "04",
        'Complete architectural freedom. “Ingenuity and a tinkerer mindset,” officially encouraged.',
        size=18,
    )


def add_challenge_row(slide, y, number, text):
    add_rect(slide, 0.76, y + 0.06, 0.035, 0.47, CORAL)
    add_text(slide, number, 0.98, y, 0.42, 0.30, 10.5, CORAL, font=CODE_FONT, bold=True)
    add_text(
        slide,
        text,
        1.48,
        y - 0.03,
        11.05,
        0.58,
        18,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_evidence_frame(slide, image_path, x, label, label_color):
    if image_path.exists():
        picture = add_picture_cover(slide, image_path, x + 0.06, 1.74, 5.75, 3.06)
        picture.name = label.replace(" ", "_")
    else:
        add_missing_asset(slide, str(image_path), x + 0.06, 1.74, 5.75, 3.06, dark=True)
    add_rect(slide, x + 0.22, 1.91, 2.18, 0.39, DEEP, radius=True, opacity=88)
    add_text(
        slide,
        label,
        x + 0.38,
        1.99,
        1.86,
        0.18,
        10.5,
        label_color,
        font=CODE_FONT,
        bold=True,
    )


def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    add_slide_title(slide, "The Challenges", "The same ground · radically different evidence", dark=True)
    add_evidence_frame(slide, SLIDE_3_SHARP, 0.72, "SHARP FRAME", GREEN)
    add_evidence_frame(slide, SLIDE_3_BLURRED, 6.74, "WORST-BLURRED FRAME", CORAL)

    add_challenge_row(
        slide,
        5.08,
        "01",
        "The field is highly repetitive — one patch of dirt looks like every other patch of dirt",
    )
    add_challenge_row(
        slide,
        5.72,
        "02",
        "Several frames are so blurred that even a human struggles to place a feature confidently",
    )
    add_challenge_row(
        slide,
        6.36,
        "03",
        "GPS telemetry on the drones is off by multiple metres — you can’t just trust the numbers you’re handed",
    )


def add_failure_mark(slide, x, y):
    first = add_rect(slide, x, y, 0.28, 0.055, CORAL)
    first.rotation = 45
    second = add_rect(slide, x, y, 0.28, 0.055, CORAL)
    second.rotation = 315
    return first, second


def add_comparison_row(slide, y, h, number, solution, reasons):
    add_rect(slide, 0.72, y, 12.0, h, WHITE, line=PALE, radius=True, line_width=1.0)
    add_rect(slide, 0.72, y, 3.24, h, MIDNIGHT, radius=True)
    add_text(
        slide,
        number,
        1.00,
        y + 0.27,
        0.38,
        0.30,
        12,
        GREEN,
        font=CODE_FONT,
        bold=True,
    )
    add_text(
        slide,
        solution,
        1.00,
        y + 0.72,
        2.55,
        0.85,
        25 if len(solution) < 14 else 21,
        WHITE,
        font=HEADER_FONT,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        "WHY IT FAILED",
        4.36,
        y + 0.30,
        2.5,
        0.28,
        10.5,
        CORAL,
        font=CODE_FONT,
        bold=True,
    )
    reason_y = y + 0.78
    for reason in reasons:
        add_failure_mark(slide, 4.38, reason_y + 0.16)
        add_text(
            slide,
            reason,
            4.88,
            reason_y,
            7.08,
            0.48,
            19,
            INK,
            font=BODY_FONT,
            valign=MSO_ANCHOR.MIDDLE,
        )
        reason_y += 0.63


def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_slide_title(slide, "What Had Failed", "Solution  |  Why it failed")
    add_comparison_row(
        slide,
        1.82,
        2.22,
        "01",
        "Homography",
        ["Lack of good joint-plane corner features", "Distracting van features"],
    )
    add_comparison_row(
        slide,
        4.36,
        2.22,
        "02",
        "GeoCalib →\nPyCeres",
        ["Lack of good distinguishable features"],
    )


def add_picker_summary(slide, x, number, text):
    add_rect(slide, x, 5.98, 0.035, 0.92, GREEN)
    add_text(
        slide,
        number,
        x + 0.20,
        5.98,
        0.38,
        0.24,
        10.5,
        GREEN,
        font=CODE_FONT,
        bold=True,
    )
    add_text(
        slide,
        text,
        x + 0.20,
        6.23,
        2.52,
        0.72,
        17,
        WHITE,
        font=BODY_FONT,
    )


def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    add_slide_title(slide, "The Manual Correspondence Picker", dark=True)

    movie = add_movie_contain(
        slide,
        SLIDE_5_VIDEO,
        SLIDE_5_POSTER,
        0.72,
        1.49,
        11.89,
        4.32,
        "Slide5_Video",
        dark=True,
    )
    if movie is not None:
        configure_video_click_timing(prs, 5, "Slide5_Video", 84467)

    add_picker_summary(slide, 0.78, "01", "Grid view supports independent zoom/pan and adjustable marker sizing.")
    add_picker_summary(slide, 3.80, "02", "Four feature types include paired van dimensions as hard constraints.")
    add_picker_summary(slide, 6.82, "03", "Marks and correspondences can be removed or browsed by feature type.")
    add_picker_summary(slide, 9.84, "04", "Unsaved changes require a second keypress before discard.")
    add_notes(
        slide,
        "The justification of this tool is to test how PyCeres performs when handed good feature matching.",
    )


def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_slide_title(slide, "The Scorer")

    if SLIDE_6_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_6_VISUAL, 0.72, 1.66, 7.16, 5.37)
        picture.name = "Scorer_Evidence"
    else:
        add_missing_asset(slide, str(SLIDE_6_VISUAL), 0.72, 1.66, 7.16, 5.37)

    add_rect(slide, 8.18, 2.45, 4.44, 2.55, MIDNIGHT, radius=True)
    add_rect(slide, 8.52, 2.92, 0.045, 1.60, CORAL)
    add_text(
        slide,
        "The scorer was built to assess manual feature-matching quality.",
        8.84,
        2.86,
        3.35,
        1.72,
        24,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )


def build_slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    add_text(
        slide,
        '“Manual Feature Matching Works Great on 7 Frames”',
        0.72,
        0.28,
        11.88,
        0.70,
        32,
        WHITE,
        font=HEADER_FONT,
        bold=True,
        valign=MSO_ANCHOR.TOP,
    )
    add_rect(slide, 0.74, 1.12, 11.85, 0.018, "344054")

    movie = add_movie_contain(
        slide,
        SLIDE_7_VIDEO,
        SLIDE_7_POSTER,
        1.59,
        1.24,
        10.15,
        5.71,
        "Slide7_Video",
        dark=True,
    )
    if movie is not None:
        configure_video_click_timing(prs, 7, "Slide7_Video", 72233)

    add_rect(slide, 1.87, 5.65, 9.59, 1.18, DEEP, radius=True, opacity=94)
    add_text(
        slide,
        "Known van geometry (wheel axis, roof edges) as hard anchors + hand-picked ground correspondences → a genuinely good calibration on the 7 clearest frames. So PyCeres is good, and our problem is the feature matching.",
        2.16,
        5.82,
        9.02,
        0.82,
        18,
        WHITE,
        font=BODY_FONT,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def build_slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    add_slide_title(slide, "The Other Six", dark=True)

    add_rect(slide, 0.76, 1.95, 0.05, 3.43, CORAL)
    add_text(
        slide,
        "Remembering past days of a base manual calibration – optimized by automatic processes, I was hoping to be able to treat the 7 manually solved frames as core, then add other frames to be solved automatically. It didn’t succeed.",
        1.04,
        1.86,
        4.66,
        3.66,
        21,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )

    if SLIDE_8_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_8_VISUAL, 6.09, 1.55, 6.52, 4.89)
        picture.name = "Trusted_Core_Failed_Extension"
    else:
        add_missing_asset(slide, str(SLIDE_8_VISUAL), 6.09, 1.55, 6.52, 4.89, dark=True)


def build_slide_9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_slide_title(slide, "The Solution")

    add_rect(slide, 0.72, 1.70, 3.34, 4.64, MIDNIGHT, radius=True)
    add_rect(slide, 1.02, 2.08, 0.045, 3.08, GREEN)
    add_text(
        slide,
        "The old stack just couldn’t crack it, and something smarter was needed. MASt3R’s dense reconstruction could handle the low-texture, repetitive ground the old feature matchers choked on.",
        1.28,
        2.04,
        2.46,
        3.82,
        18,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.TOP,
    )
    if SLIDE_9_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_9_VISUAL, 4.34, 1.70, 8.27, 4.65)
        picture.name = "Sparse_to_Dense"
    else:
        add_missing_asset(slide, str(SLIDE_9_VISUAL), 4.34, 1.70, 8.27, 4.65)
    add_notes(slide, "Mention that this magic requires better hardware.")


def build_slide_10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_slide_title(slide, "Step Zero: Get a GPU")

    if SLIDE_10_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_10_VISUAL, 0.72, 1.68, 7.56, 4.26)
        picture.name = "GPU_Provider_Decision"
    else:
        add_missing_asset(slide, str(SLIDE_10_VISUAL), 0.72, 1.68, 7.56, 4.26)

    add_rect(slide, 8.56, 1.68, 4.05, 4.92, MIDNIGHT, radius=True)
    add_rect(slide, 8.88, 2.08, 0.045, 3.82, GREEN)
    add_text(
        slide,
        "RunPod, chosen over Google Colab specifically because Colab’s runtimes wipe on timeout — meaning MASt3R’s CUDA/C++ kernels would need recompiling from scratch every session, versus RunPod’s persistent volume storage compiling once and staying ready.",
        9.18,
        2.02,
        3.04,
        3.96,
        18,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )


def build_slide_11(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    add_slide_title(slide, "Compute Headless in the Cloud, Test Locally", dark=True)

    add_rect(slide, 0.76, 1.78, 0.045, 4.58, GREEN)
    add_rich_line(
        slide,
        [
            {"text": "RunPod has no display; the interactive viewer needs one. ", "color": WHITE},
            {"text": "--export-solve", "font": CODE_FONT, "bold": True, "color": GREEN},
            {"text": " runs the full pipeline remotely and serializes camera poses + terrain point clouds to disk. Those files travel to the local machine, where ", "color": WHITE},
            {"text": "--import-solve", "font": CODE_FONT, "bold": True, "color": GREEN},
            {"text": " skips OCR/MASt3R/Ceres entirely and opens the viewer straight against the already-solved cameras.", "color": WHITE},
        ],
        1.06,
        1.73,
        3.35,
        4.72,
        18,
        color=WHITE,
    )
    if SLIDE_11_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_11_VISUAL, 4.63, 1.60, 7.98, 4.50)
        picture.name = "Cloud_to_Local_Pipeline"
    else:
        add_missing_asset(slide, str(SLIDE_11_VISUAL), 4.63, 1.60, 7.98, 4.50, dark=True)


def build_slide_12(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_slide_title(slide, "MASt3R’s Achilles’ Heel")

    if SLIDE_12_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_12_VISUAL, 0.72, 1.60, 8.86, 5.01)
        picture.name = "Drift_to_Anchor"
    else:
        add_missing_asset(slide, str(SLIDE_12_VISUAL), 0.72, 1.60, 8.86, 5.01)

    add_rect(slide, 9.82, 2.17, 2.79, 3.52, MIDNIGHT, radius=True)
    add_rect(slide, 10.13, 2.64, 0.045, 2.56, CORAL)
    add_text(
        slide,
        "MASt3R has one weakness — it drifts. Qwen finds a shared anchor in the images to constrain that drift.",
        10.43,
        2.57,
        1.83,
        2.72,
        21,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )


def build_slide_13(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_text(
        slide,
        "Qwen’s Job: Find the One Thing Every Frame Shares",
        0.72,
        0.40,
        11.88,
        0.70,
        32,
        INK,
        font=HEADER_FONT,
        bold=True,
    )
    add_rect(slide, 0.74, 1.39, 11.85, 0.018, PALE)

    if SLIDE_13_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_13_VISUAL, 0.72, 1.67, 8.18, 4.09)
        picture.name = "Qwen_Anchor_Discovery"
    else:
        add_missing_asset(slide, str(SLIDE_13_VISUAL), 0.72, 1.67, 8.18, 4.09)

    add_rect(slide, 9.16, 1.64, 3.45, 5.37, MIDNIGHT, radius=True)
    add_rect(slide, 9.47, 2.04, 0.045, 4.53, GREEN)
    add_text(
        slide,
        "Per frame: list every man-made object, bbox + label. Across all frames: consolidate labels, rank by how many frames contain each one. Among the top-coverage candidates, a second Qwen call — reasoning over label text alone — picks the single best fixed 3D reference point. Output: that object’s pixel centroid in every frame it appears in, weighted by CLIP confidence.",
        9.78,
        1.98,
        2.46,
        4.66,
        17.5,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )


def build_slide_14(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    add_rich_line(
        slide,
        [
            {"text": "--export-mesh", "font": CODE_FONT, "bold": True, "color": GREEN, "size": 34},
            {"text": ": Visual Debugging", "font": HEADER_FONT, "bold": True, "color": WHITE, "size": 38},
        ],
        0.72,
        0.42,
        11.85,
        0.62,
        38,
        color=WHITE,
    )
    add_rect(slide, 0.74, 1.39, 11.85, 0.018, "344054")

    add_rect(slide, 0.72, 1.68, 3.54, 4.98, DEEP, radius=True)
    add_rect(slide, 1.02, 2.07, 0.045, 3.82, GREEN)
    add_text(
        slide,
        "Even with the right tools, the reconstruction wasn’t good. So it was time to see things with our eyes: a debug tool that dumps the raw MASt3R reconstruction + solved cameras (and, once solved, the Ceres-refined + Sim(3)-aligned version) as real .glb scenes, importable straight into Blender.",
        1.32,
        2.00,
        2.56,
        3.98,
        18,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )
    if SLIDE_14_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_14_VISUAL, 4.53, 1.68, 8.08, 4.55)
        picture.name = "Blender_Debug_Evidence"
    else:
        add_missing_asset(slide, str(SLIDE_14_VISUAL), 4.53, 1.68, 8.08, 4.55, dark=True)


def add_attention_frame(slide, image_path, x, label, accent):
    if image_path.exists():
        picture = add_picture_cover(slide, image_path, x, 1.78, 5.88, 3.31)
        picture.name = label.replace(" ", "_")
    else:
        add_missing_asset(slide, str(image_path), x, 1.78, 5.88, 3.31, dark=True)
    add_rect(slide, x + 0.20, 1.99, 2.30, 0.40, DEEP, radius=True, opacity=88)
    add_text(
        slide,
        label,
        x + 0.37,
        2.07,
        1.96,
        0.18,
        10.5,
        accent,
        font=CODE_FONT,
        bold=True,
    )


def build_slide_15(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    add_slide_title(slide, "Attention Magnets", dark=True)
    add_attention_frame(slide, SLIDE_15_BLACK, 0.72, "HARD BLACK FILL", CORAL)
    add_attention_frame(slide, SLIDE_15_FEATHERED, 6.73, "FEATHERED GREY FILL", GREEN)

    add_rect(slide, 0.92, 5.45, 0.045, 1.04, CORAL)
    add_text(
        slide,
        'Hard black HUD fills acted as infinite-contrast “attention magnets” for MASt3R’s vision transformer, warping its initial pose estimates. Fixed with a feathered gray fill — Gaussian-blurred, no hard edge.',
        1.22,
        5.35,
        11.05,
        1.20,
        20,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
        align=PP_ALIGN.CENTER,
    )


def build_slide_16(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_slide_title(slide, "Flat Earth")

    add_rect(slide, 0.72, 1.68, 3.58, 4.98, MIDNIGHT, radius=True)
    add_rect(slide, 1.03, 2.08, 0.045, 3.82, GREEN)
    add_text(
        slide,
        "Replaced the flat z = 0 ground-plane assumption with an interpolated height field built from the solved 3D points themselves — ray-casting now iterates to find where a ray actually meets the real, bumpy terrain, falling back to flat only if that doesn’t converge.",
        1.33,
        2.00,
        2.59,
        3.98,
        18,
        WHITE,
        font=BODY_FONT,
        valign=MSO_ANCHOR.MIDDLE,
    )
    if SLIDE_16_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_16_VISUAL, 4.56, 1.68, 8.05, 4.53)
        picture.name = "Terrain_Aware_Raycast"
    else:
        add_missing_asset(slide, str(SLIDE_16_VISUAL), 4.56, 1.68, 8.05, 4.53)


def build_slide_17(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    movie = add_movie_contain(
        slide,
        SLIDE_17_VIDEO,
        SLIDE_17_POSTER,
        0,
        0,
        SLIDE_W,
        SLIDE_H,
        "Slide17_Video",
        dark=True,
    )
    panel = add_rect(slide, 0.48, 0.38, 4.58, 0.95, DEEP, radius=True, opacity=75)
    panel.name = "Slide17_TitlePanel"
    title = add_text(
        slide,
        "See It In Action",
        0.78,
        0.55,
        3.98,
        0.54,
        38,
        WHITE,
        font=HEADER_FONT,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    title.name = "Slide17_TitleText"
    set_text_outline(title, "000000", 0.5)
    if movie is not None:
        configure_video_click_timing(
            prs,
            17,
            "Slide17_Video",
            119533,
            fade_shape_names=["Slide17_TitlePanel", "Slide17_TitleText"],
            fade_duration_ms=350,
        )


def add_result_item(slide, y, lead, label, detail=None, accent=GREEN):
    add_rect(slide, 8.63, y, 0.055, 0.70, accent)
    add_text(
        slide,
        lead,
        8.91,
        y - 0.03,
        3.66,
        0.38,
        23,
        WHITE,
        font=HEADER_FONT,
        bold=True,
    )
    add_text(
        slide,
        label,
        8.91,
        y + 0.36,
        3.66,
        0.30,
        11.5,
        PALE,
        font=CODE_FONT,
        bold=True,
    )
    if detail:
        add_text(
            slide,
            detail,
            8.91,
            y + 0.70,
            3.66,
            0.38,
            13.5,
            PALE,
            font=BODY_FONT,
        )


def build_slide_18(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_text(
        slide,
        "How the Result Was Measured",
        0.72,
        0.14,
        11.85,
        0.52,
        34,
        INK,
        font=HEADER_FONT,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_rect(slide, 0.74, 0.75, 11.85, 0.018, PALE)

    if SLIDE_18_VISUAL.exists():
        picture = add_picture_contain(slide, SLIDE_18_VISUAL, 1.75, 0.83, 9.83, 5.35)
        picture.name = "Measurement_Method_Evidence"
    else:
        add_missing_asset(slide, str(SLIDE_18_VISUAL), 1.75, 0.83, 9.83, 5.35)

    add_rect(slide, 0.45, 6.23, 12.43, 1.07, MIDNIGHT)
    methodology = [
        ("01", "7 ground features, manually marked in up to 13 frames: 62 judged marks in total"),
        ("02", "Every marked frame was tested as the starting view against every other marked frame"),
        ("03", "650 directed source-to-target cross-checks: 649 projected into-frame, 1 projected out-of-frame, 0 ray misses"),
        ("04", "The score is the distance between the system's projected point and the manually marked location"),
        ("05", "Metres are the headline metric because pixel error favors distant, zoomed-out frames"),
    ]
    column_width = 12.43 / len(methodology)
    for index, (number, item_text) in enumerate(methodology):
        column_x = 0.45 + index * column_width
        add_text(slide, number, column_x + 0.15, 6.34, 0.42, 0.22, 10.0, GREEN, font=CODE_FONT, bold=True)
        add_text(slide, item_text, column_x + 0.15, 6.56, column_width - 0.30, 0.63, 9.5, WHITE)
        if index < len(methodology) - 1:
            add_rect(slide, column_x + column_width - 0.012, 6.38, 0.012, 0.73, "455166")

    add_notes(
        slide,
        "The manual marks are imperfect, especially in blurred distant frames. Errors below roughly one metre can be difficult to distinguish from click uncertainty. The result should therefore be presented as a measured engineering evaluation, not false precision.",
    )


def build_slide_19(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, OFFWHITE)
    add_slide_title(
        slide,
        "Measured Result: About One Van Length",
        "650 cross-checks across all 13 frames",
    )

    add_rect(slide, 0.54, 1.66, 7.77, 5.35, WHITE, line=PALE, radius=True, line_width=1.0)
    matrix = add_picture_contain(slide, RESULTS_MATRIX, 0.65, 1.76, 7.55, 4.73)
    matrix.name = "CrossReprojection_Matrix"
    add_text(
        slide,
        "The system consistently finds the correct area of the field across extremely difficult imagery, with a stable measured error and a visible frame-by-frame failure pattern.",
        0.82,
        6.49,
        7.20,
        0.43,
        10.6,
        MUTED,
        italic=True,
        align=PP_ALIGN.CENTER,
    )

    add_rect(slide, 8.43, 1.66, 4.36, 5.35, MIDNIGHT, radius=True)
    add_rect(slide, 8.68, 1.92, 0.055, 0.70, GREEN)
    add_text(slide, "4.05 m median view error — roughly one van length", 8.94, 1.88, 3.45, 0.72, 19.5, WHITE, font=HEADER_FONT, bold=True)
    add_rect(slide, 8.68, 2.74, 3.73, 0.012, "455166")
    add_text(slide, "38.9 px median image error — 48 of 649 in-frame checks, about 7%, landed within 10 px", 8.94, 2.88, 3.45, 0.82, 16.0, WHITE)
    add_rect(slide, 8.68, 3.80, 3.73, 0.012, "455166")
    add_text(slide, "Repeatable: approximately 4.2 m in the first picking session and 4.0 m in the fresh session", 8.94, 3.94, 3.45, 0.82, 16.0, WHITE)
    add_rect(slide, 8.68, 4.84, 3.73, 0.012, "455166")
    add_text(slide, "Best source: 04709 at 2.67 m; worst source: 04569 at 8.08 m", 8.94, 4.98, 3.45, 0.64, 15.5, WHITE, bold=True)
    add_rect(slide, 8.68, 5.70, 3.73, 0.012, CORAL)
    add_text(
        slide,
        "Frame 05934 has a physically invalid camera solution below the ground with flipped orientation. Its anchor is usable, but its camera axes are not.",
        8.94,
        5.86,
        3.45,
        0.88,
        12.1,
        CORAL,
        bold=True,
    )
    add_notes(
        slide,
        "Pixel error varies strongly with distance and framing, so the metre-based view error is the fairer headline comparison. 12035 is one of the best frames to start from but one of the hardest to land in because its camera looks down much more steeply than the others. The result is not the original ≤10 px target across the dataset; it is a stable, quantified solution over all 13 hostile frames.",
    )


def build_appendix_a1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    add_slide_title(slide, "Better Is the Enemy of Good", dark=True)

    add_rect(slide, 0.72, 1.68, 5.05, 5.12, "172235", line="344054", radius=True, line_width=1.0)
    add_rect(slide, 0.98, 1.95, 0.06, 4.56, GREEN)
    add_text(
        slide,
        "An incremental two-stage Ceres solve (lock down the good frames first, register the blurry ones motion-only against that structure). Then, one level deeper: splitting MASt3R's own reconstruction the same way — a clean backbone built only from good frames, with each bad frame individually resectioned against it so it can never contaminate another bad frame or the backbone itself. Unfortunately, neither produced better real-world output.",
        1.30,
        1.96,
        4.10,
        4.50,
        18.0,
        WHITE,
    )

    add_rect(slide, 6.02, 1.68, 6.77, 4.96, DEEP, line="344054", radius=True, line_width=1.0)
    if APPENDIX_A1_VISUAL.exists():
        picture = add_picture_contain(slide, APPENDIX_A1_VISUAL, 6.14, 1.80, 6.53, 4.72)
        picture.name = "Appendix_Ceres_Reconstruction"
    else:
        add_missing_asset(slide, str(APPENDIX_A1_VISUAL), 6.14, 1.80, 6.53, 4.72, dark=True)

    add_notes(
        slide,
        "Worth mentioning that this whole effort was across half a day. As feedback was already good, we decided to move on to the next challenge, knowing we can continue fighting it another day.",
    )


def build_slide_20(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MIDNIGHT)
    add_rect(slide, 4.87, 2.18, 3.60, 0.055, GREEN)
    add_text(
        slide,
        "Thank You. Questions?",
        1.00,
        2.63,
        11.33,
        1.05,
        54,
        WHITE,
        font=HEADER_FONT,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_rect(slide, 5.99, 4.17, 1.35, 0.055, GREEN)


SLIDE_BUILDERS = {
    "1": build_slide_1,
    "2": build_slide_2,
    "3": build_slide_3,
    "4": build_slide_4,
    "5": build_slide_5,
    "6": build_slide_6,
    "7": build_slide_7,
    "8": build_slide_8,
    "9": build_slide_9,
    "10": build_slide_10,
    "11": build_slide_11,
    "12": build_slide_12,
    "13": build_slide_13,
    "14": build_slide_14,
    "15": build_slide_15,
    "16": build_slide_16,
    "17": build_slide_17,
    "18": build_slide_18,
    "19": build_slide_19,
    "A1": build_appendix_a1,
    "20": build_slide_20,
}

PROOF_SLIDES = ["1", "4", "19"]
BATCH_SLIDES = {
    1: ["1", "2", "3", "4"],
    2: ["1", "2", "3", "4", "5", "6", "7", "8"],
    3: ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12", "13"],
    4: ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12", "13", "14", "15", "16", "17"],
    5: ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12", "13", "14", "15", "16", "17", "18", "19", "A1", "20"],
}


def build_deck(output_path, slide_ids, title, subject):
    WARNINGS.clear()
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    for slide_id in slide_ids:
        SLIDE_BUILDERS[slide_id](prs)
    prs.core_properties.title = title
    prs.core_properties.subject = subject
    prs.core_properties.author = "Raycast Challenge"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return list(WARNINGS)


def main():
    parser = argparse.ArgumentParser(description="Build the Raycast Challenge v3 deck.")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--proof", action="store_true", help="Build the locked Slides 1, 4, and 19 proof.")
    mode.add_argument("--through-batch", type=int, choices=sorted(BATCH_SLIDES), metavar="N")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    if args.proof:
        slide_ids = PROOF_SLIDES
        title = "Raycast Challenge v3 — Design Proof"
        subject = "Slides 1, 4, and 19"
    else:
        slide_ids = BATCH_SLIDES[args.through_batch]
        title = f"Raycast Challenge v3 — Through Batch {args.through_batch}"
        subjects = {
            1: "Challenge: Slides 1–4",
            2: "Challenge and Manual Baseline: Slides 1–8",
            3: "Challenge, Manual Baseline, and New Architecture: Slides 1–13",
            4: "Challenge through Debugging and Demo: Slides 1–17",
            5: "Complete Raycast Challenge v3: Slides 1–19, Appendix A1, and Slide 20",
        }
        subject = subjects[args.through_batch]

    warnings = build_deck(args.output.resolve(), slide_ids, title, subject)
    print(f"Wrote {args.output.resolve()}")
    print(f"Slides: {', '.join(slide_ids)}")
    for warning in warnings:
        print(f"WARNING: {warning}")


if __name__ == "__main__":
    main()
