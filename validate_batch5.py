import argparse
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


parser = argparse.ArgumentParser(description="Validate a Raycast Challenge v3 deck.")
parser.add_argument(
    "path",
    nargs="?",
    type=Path,
    default=Path(__file__).resolve().parent / "presentation_v3.pptx",
)
parser.add_argument(
    "--linked",
    action="store_true",
    help="Require external relative video links and no embedded MP4 payloads.",
)
args = parser.parse_args()
path = args.path.resolve()
prs = Presentation(path)
assert len(prs.slides) == 21

expected = {
    18: [
        "How the Result Was Measured",
        "7 ground features, manually marked in up to 13 frames: 62 judged marks in total",
        "Every marked frame was tested as the starting view against every other marked frame",
        "650 directed source-to-target cross-checks: 649 projected into-frame, 1 projected out-of-frame, 0 ray misses",
        "The score is the distance between the system's projected point and the manually marked location",
        "Metres are the headline metric because pixel error favors distant, zoomed-out frames",
    ],
    19: [
        "Measured Result: About One Van Length",
        "650 CROSS-CHECKS ACROSS ALL 13 FRAMES",
        "4.05 m median view error — roughly one van length",
        "38.9 px median image error — 48 of 649 in-frame checks, about 7%, landed within 10 px",
        "Repeatable: approximately 4.2 m in the first picking session and 4.0 m in the fresh session",
        "Best source: 04709 at 2.67 m; worst source: 04569 at 8.08 m",
        "The system consistently finds the correct area of the field across extremely difficult imagery, with a stable measured error and a visible frame-by-frame failure pattern.",
        "Frame 05934 has a physically invalid camera solution below the ground with flipped orientation. Its anchor is usable, but its camera axes are not.",
    ],
    20: [
        "Better Is the Enemy of Good",
        "An incremental two-stage Ceres solve (lock down the good frames first, register the blurry ones motion-only against that structure). Then, one level deeper: splitting MASt3R's own reconstruction the same way — a clean backbone built only from good frames, with each bad frame individually resectioned against it so it can never contaminate another bad frame or the backbone itself. Unfortunately, neither produced better real-world output.",
    ],
    21: ["Thank You. Questions?"],
}

for slide_number, required in expected.items():
    slide_text = "\n".join(
        shape.text
        for shape in prs.slides[slide_number - 1].shapes
        if getattr(shape, "has_text_frame", False)
    )
    missing = [item for item in required if item not in slide_text]
    assert not missing, (slide_number, missing)

expected_notes = {
    18: "The manual marks are imperfect, especially in blurred distant frames. Errors below roughly one metre can be difficult to distinguish from click uncertainty. The result should therefore be presented as a measured engineering evaluation, not false precision.",
    19: "Pixel error varies strongly with distance and framing, so the metre-based view error is the fairer headline comparison. 12035 is one of the best frames to start from but one of the hardest to land in because its camera looks down much more steeply than the others. The result is not the original ≤10 px target across the dataset; it is a stable, quantified solution over all 13 hostile frames.",
    20: "Worth mentioning that this whole effort was across half a day. As feedback was already good, we decided to move on to the next challenge, knowing we can continue fighting it another day.",
}

for slide_number, note in expected_notes.items():
    actual = prs.slides[slide_number - 1].notes_slide.notes_text_frame.text
    assert actual == note, (slide_number, actual)

media_shapes = []
for number, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.MEDIA:
            media_shapes.append((number, shape.name))
assert media_shapes == [
    (5, "Slide5_Video"),
    (7, "Slide7_Video"),
    (17, "Slide17_Video"),
], media_shapes

with ZipFile(path) as archive:
    names = archive.namelist()
    media = sorted(name for name in names if name.startswith("ppt/media/"))
    video_parts = [name for name in media if name.lower().endswith(".mp4")]
    slide17 = archive.read("ppt/slides/slide17.xml").decode("utf-8")
    assert "presetClass=\"exit\"" in slide17
    assert slide17.count("filter=\"fade\"") == 2
    assert "dur=\"350\"" in slide17
    assert all(
        name in names
        for name in [
            "ppt/slides/slide18.xml",
            "ppt/slides/slide19.xml",
            "ppt/slides/slide20.xml",
            "ppt/slides/slide21.xml",
        ]
    )
    if args.linked:
        assert not video_parts, video_parts
        linked_videos = {
            5: "visuals/Raycast_Slide_8_Visual_RETIMED.mp4",
            7: "visuals/Raycast_Slide_10_Visual.mp4",
            17: "visuals/Raycast_Slide_35_Visual.mp4",
        }
        for slide_number, target in linked_videos.items():
            rels = archive.read(
                f"ppt/slides/_rels/slide{slide_number}.xml.rels"
            ).decode("utf-8")
            slide_xml = archive.read(f"ppt/slides/slide{slide_number}.xml").decode(
                "utf-8"
            )
            assert rels.count(f'Target="{target}"') == 2
            assert rels.count('TargetMode="External"') == 2
            assert "r:link=" in slide_xml and "p14:media" in slide_xml
    else:
        assert len(video_parts) == 3, video_parts

media_mode = "linked" if args.linked else "embedded"
print(
    "Validation passed: 21 slides; Batch 5 text and notes exact; "
    f"3 {media_mode} videos; Slide 17 has two 350 ms exit fades."
)
print(f"Media parts: {len(media)}")
